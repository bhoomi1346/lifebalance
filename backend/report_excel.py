import io
from datetime import date, timedelta

from flask import Blueprint, request, send_file, jsonify
from flask_jwt_extended import jwt_required, get_jwt_identity
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from models import DailyLog, Prediction
from reports import period_for, summarise, compare, MIN_DAYS, DISCLAIMER

excel_bp = Blueprint("report_excel", __name__)

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1E6F5C")


def latest_prediction(log_id):
    return (Prediction.query.filter_by(log_id=log_id)
            .order_by(Prediction.created_at.desc()).first())


def write_header(ws, row, headers):
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_summary_sheet(wb, kind, start, end, prev_start, prev_end, message, current, previous, cmp):
    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 22

    ws["A1"] = "Life Balance Report"
    ws["A1"].font = Font(bold=True, size=14, color="1E6F5C")
    rows = [
        ("Report type", kind.capitalize()),
        ("Period", f"{start.isoformat()} to {end.isoformat()}"),
        ("Previous period", f"{prev_start.isoformat()} to {prev_end.isoformat()}"),
        ("Generated", date.today().isoformat()),
        ("Logged days", current["days"]),
        ("Average sleep (h)", current["avg_sleep"]),
        ("Average working hours (h)", current["avg_work"]),
        ("Average activity (min per day)", current["avg_activity"]),
        ("Average balance score (0-100)", current["avg_wlb"]),
        ("Low stress days", current["stress_counts"]["Low"]),
        ("Moderate stress days", current["stress_counts"]["Moderate"]),
        ("High stress days", current["stress_counts"]["High"]),
        ("Change in sleep (h)", cmp["sleep_change"]),
        ("Change in working hours (h)", cmp["work_change"]),
        ("Change in activity (min)", cmp["activity_change"]),
        ("Change in balance score", cmp["wlb_change"]),
        ("Change in high-stress days", cmp["high_days_change"]),
        ("Best day", f"{current['best_day']['date']} (score {current['best_day']['wlb_score']})"
                     if current["best_day"] else "No scored days"),
        ("Most difficult day", f"{current['hardest_day']['date']} (score {current['hardest_day']['wlb_score']})"
                               if current["hardest_day"] else "No scored days"),
    ]
    r = 3
    for label, value in rows:
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=2, value="No data" if value is None else value)
        r += 1

    if message:
        r += 1
        ws.cell(row=r, column=1, value=message).alignment = Alignment(wrap_text=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.row_dimensions[r].height = 45

    r += 2
    ws.cell(row=r, column=1, value="Most frequent suggestions").font = Font(bold=True)
    r += 1
    if current["top_tips"]:
        for item in current["top_tips"]:
            ws.cell(row=r, column=1, value=item["tip"]).alignment = Alignment(wrap_text=True)
            ws.cell(row=r, column=2, value=f"given {item['times']} times")
            ws.row_dimensions[r].height = 30
            r += 1
    else:
        ws.cell(row=r, column=1, value="No suggestions saved in this period.")
        r += 1

    r += 1
    ws.cell(row=r, column=1, value=DISCLAIMER).font = Font(italic=True, color="5B6B68")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)


def build_raw_sheet(wb, user_id, start, end):
    ws = wb.create_sheet("Raw data")
    headers = ["Date", "Sleep (h)", "Working hours", "Activity (min)", "Water (L)",
               "Meals", "Mood", "Self-reported stress", "Stress estimate",
               "Balance score", "Balance band"]
    write_header(ws, 1, headers)
    widths = [12, 11, 14, 15, 11, 20, 8, 20, 17, 15, 15]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    logs = (DailyLog.query.filter(DailyLog.user_id == user_id,
                                  DailyLog.log_date >= start,
                                  DailyLog.log_date <= end)
            .order_by(DailyLog.log_date).all())
    for r, lg in enumerate(logs, start=2):
        p = latest_prediction(lg.log_id)
        ws.append([
            lg.log_date.isoformat(), lg.sleep_hours, lg.working_hours, lg.activity_minutes,
            lg.water_litres, lg.meal_regularity, lg.mood, lg.self_reported_stress or "",
            p.stress_level if p else "Not scored",
            p.wlb_score if p else "",
            p.wlb_band if p else "",
        ])
    ws.freeze_panes = "A2"


@excel_bp.route("/excel", methods=["GET"])
@jwt_required()
def report_excel():
    user_id = int(get_jwt_identity())
    kind = request.args.get("type", "weekly")
    today = date.today()

    start, end, error = period_for(kind, today, request.args.get("from"), request.args.get("to"))
    if error:
        return jsonify({"error": error}), 400

    length = (end - start).days + 1
    prev_end = start - timedelta(days=1)
    prev_start = prev_end - timedelta(days=length - 1)

    current = summarise(user_id, start, end)
    previous = summarise(user_id, prev_start, prev_end)
    cmp = compare(current, previous)

    needed = MIN_DAYS.get(kind, 3)
    if current["days"] == 0:
        message = "There are no records in this period yet. Add daily entries to build a report."
    elif current["days"] < needed:
        message = (f"You have {current['days']} logged day(s) in this period. A full {kind} report "
                   f"needs at least {needed}. The available days are shown below.")
    else:
        message = None

    wb = Workbook()
    build_summary_sheet(wb, kind, start, end, prev_start, prev_end, message, current, previous, cmp)
    build_raw_sheet(wb, user_id, start, end)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"life_balance_{kind}_{start.isoformat()}_{end.isoformat()}.xlsx"
    return send_file(buf,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name=filename)